"""End-to-end tests: index → hybrid search → API citations."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch

import numpy as np
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from holusight import config as config_module
from holusight.api import Holusight
from holusight.config import ServerConfig
from holusight.search import rrf_merge

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "pilot_docs"


@pytest.fixture
def indexed_engine(tmp_path, monkeypatch):
    """Index pilot fixtures into an isolated data directory."""
    monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "data")
    engine = Holusight(FIXTURES, config=ServerConfig())
    stats = engine.index(force_rebuild=True)
    assert stats.files_indexed >= 2
    assert stats.total_chunks >= 2
    return engine


@pytest.fixture
def pptx_engine(tmp_path, monkeypatch):
    """Use only generated slides, a private index, and deterministic local vectors."""
    class FakeEmbedder:
        def embed(self, texts):
            return np.tile(np.array([1.0, 0.0], dtype=np.float32), (len(texts), 1))

        def embed_query(self, text):
            return self.embed([text])[0]

    embedder = FakeEmbedder()
    monkeypatch.setenv("HF_HUB_OFFLINE", "1")
    monkeypatch.setenv("TRANSFORMERS_OFFLINE", "1")
    for key in ("VOYAGE_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr("holusight.indexer.VOYAGE_API_KEY", None)
    monkeypatch.setattr("holusight.search.VOYAGE_API_KEY", None)
    monkeypatch.setattr("holusight.api.get_embedder", lambda *a, **kw: embedder)
    monkeypatch.setattr("holusight.indexer.get_embedder", lambda *a, **kw: embedder)
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    engine = Holusight(corpus, config=ServerConfig(
        embedding_model="synthetic-pptx", embedding_backend="local", embedding_dim=2,
        reranker=False, query_enhancement=False, metadata_boost=False, cnfb_alpha=0.0,
    ))
    yield engine
    if engine._store is not None:
        engine.store.close()


class TestPptxE2E:
    @pytest.mark.parametrize("placeholder", [True, False], ids=["placeholder", "textbox"])
    def test_visible_text_is_searchable(self, pptx_engine, placeholder):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1 if placeholder else 6])
        shape = (slide.placeholders[1] if placeholder else slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(3), Inches(1),
        ))
        assert shape.is_placeholder is placeholder
        shape.text = "Textboxneedle visible slide content"
        path = pptx_engine.folder_path / "generated.pptx"
        prs.save(path)
        original = path.read_bytes()

        stats = pptx_engine.index()
        results = pptx_engine.search("Textboxneedle", file_glob="*.pptx")

        assert stats.files_indexed == 1
        assert stats.total_chunks == 1
        assert len(results) == 1
        result = results[0]
        assert result.snippet == shape.text
        assert result.file_path == path.name
        assert result.start_line == result.end_line == 1
        assert result.scope == "page 1"
        assert path.read_bytes() == original
        assert list(pptx_engine.folder_path.iterdir()) == [path]

    def test_decorative_shapes_preserve_later_slide_citations(self, pptx_engine):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])  # Empty slide must not renumber citations.
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        decoration = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(1),
        )
        assert decoration.is_placeholder is False
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text = "Decoratedneedle visible textbox"
        later = prs.slides.add_slide(prs.slide_layouts[1])
        later.shapes.title.text = "Laterneedle title"
        later.placeholders[1].text = "Laterbodyneedle visible body"
        path = pptx_engine.folder_path / "mixed.pptx"
        prs.save(path)
        original = path.read_bytes()

        stats = pptx_engine.index()
        assert stats.total_chunks == 2
        for query, number, scope, text in (
            ("Decoratedneedle", 2, "page 2", box.text),
            ("Laterneedle", 3, later.shapes.title.text,
             "Laterneedle title\nLaterbodyneedle visible body"),
        ):
            results = pptx_engine.search(query, file_glob="*.pptx", top_k=2)
            matching = [result for result in results if query in result.snippet]
            assert len(matching) == 1
            result = matching[0]
            assert result.file_path == path.name
            assert result.start_line == result.end_line == number
            assert result.scope == scope
            assert result.snippet == text
        assert path.read_bytes() == original
        assert list(pptx_engine.folder_path.iterdir()) == [path]


@pytest.mark.parametrize("edit_second_chunk", [False, True], ids=["unchanged", "partial-edit"])
def test_incremental_partial_file_matches_fresh_index(tmp_path, monkeypatch, edit_second_chunk):
    """An edit must not erase a sibling chunk that hash deduplication skips."""
    import numpy as np

    embedded = []

    class SyntheticEmbedder:
        def embed(self, texts):
            embedded.extend(texts)
            return np.array([[1.0, float("changed" in text)] for text in texts], dtype=np.float32)

        def embed_query(self, text):
            return np.array([1.0, 0.0], dtype=np.float32)

    embedder = SyntheticEmbedder()
    monkeypatch.setattr("holusight.indexer.get_embedder", lambda *a, **kw: embedder)
    monkeypatch.setattr("holusight.api.get_embedder", lambda *a, **kw: embedder)
    monkeypatch.setattr("holusight.indexer.VOYAGE_API_KEY", None)
    monkeypatch.setattr("holusight.search.VOYAGE_API_KEY", None)
    monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "data")
    config = ServerConfig(
        embedding_model="synthetic", embedding_dim=2,
        chunk_max_lines=2, chunk_overlap_lines=0,
        reranker=False, query_enhancement=False, metadata_boost=False, cnfb_alpha=0,
    )
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    source = corpus / "sample.txt"
    original = "alpha\none\nbeta\ntwo"
    source.write_text(original)
    engine = Holusight(corpus, config=config)
    initial = engine.index()
    assert initial.total_chunks == initial.chunks_created == 2
    before = engine.search("alpha", top_k=10)
    retained = next(result for result in before if "alpha" in result.snippet)
    retained_metadata = engine.store.get_chunk_metadata([retained.chunk_id])
    retained_vectors = engine.store.get_chunk_vectors([retained.chunk_id])
    assert len(retained_vectors) == 1
    assert len(embedded) == 2
    engine.store.close()

    edited = original.replace("two", "changed") if edit_second_chunk else original
    source.write_text(edited)
    embedded.clear()
    incremental = engine.index()
    actual = engine.search("alpha", top_k=10)
    incremental_texts = list(embedded)

    # Counterfactual: an independent fresh index of the identical edited corpus.
    fresh_corpus = tmp_path / "fresh-corpus"
    fresh_corpus.mkdir()
    (fresh_corpus / "sample.txt").write_text(edited)
    fresh_engine = Holusight(fresh_corpus, config=config)
    try:
        fresh = fresh_engine.index()
        expected = fresh_engine.search("alpha", top_k=10)
        assert fresh.total_chunks == 2
        assert {(r.start_line, r.end_line) for r in expected} == {(1, 2), (3, 4)}
        assert len(expected) == 2
        assert incremental.chunks_created == int(edit_second_chunk)
        assert incremental.chunks_skipped_unchanged == 2 - int(edit_second_chunk)
        assert len(incremental_texts) == int(edit_second_chunk)
        assert all("changed" in text and "alpha" not in text for text in incremental_texts)
        assert len(actual) == len(expected), "Partial edit lost a searchable unchanged chunk"
        assert incremental.total_chunks == fresh.total_chunks
        assert {r.chunk_id for r in actual} == {r.chunk_id for r in expected}
        assert {(r.file_path, r.start_line, r.end_line, r.snippet) for r in actual} == {
            (r.file_path, r.start_line, r.end_line, r.snippet) for r in expected
        }
        assert engine.store.get_chunk_metadata([retained.chunk_id]) == retained_metadata
        np.testing.assert_array_equal(
            engine.store.get_chunk_vectors([retained.chunk_id]), retained_vectors,
        )
        assert engine.store.bm25_search("alpha") == [retained.chunk_id]
        assert source.read_text() == edited

        engine.store.close()
        embedded.clear()
        repeat = engine.index()
        assert repeat.total_chunks == repeat.chunks_skipped_unchanged == 2
        assert repeat.chunks_created == 0
        assert embedded == []
    finally:
        engine.store.close()
        fresh_engine.store.close()


class TestE2ERetrieval:
    def test_bm25_finds_exact_payment_terms(self, indexed_engine):
        results = indexed_engine.search("Net 30 payment terms")
        assert results, "BM25 should match exact phrase in payment-terms.md"
        paths = {r.file_path for r in results}
        assert any("payment-terms" in p for p in paths)

    def test_semantic_search_finds_auth_topic(self, indexed_engine):
        results = indexed_engine.search("validate bearer tokens")
        assert results, "Vector search should find auth_utils.py"
        assert any("auth_utils" in r.file_path for r in results)

    def test_rrf_merges_keyword_and_semantic(self, indexed_engine):
        vector_ids = ["shared", "vector-only"]
        bm25_ids = ["bm25-only", "shared"]
        merged = rrf_merge([vector_ids, bm25_ids])
        assert [chunk_id for chunk_id, _ in merged] == ["shared", "bm25-only", "vector-only"]
        assert merged[0][1] == (1 / 61) + (1 / 62)
        assert merged[0][1] > merged[1][1]

        results = indexed_engine.search("payment billing invoice")
        assert results
        top = results[0]
        assert top.file_path
        assert top.start_line >= 1
        assert top.end_line >= top.start_line
        assert top.snippet
        assert top.chunk_id

    def test_citation_metadata_for_markdown(self, indexed_engine):
        results = indexed_engine.search("Net 30", file_glob="*.md")
        assert results
        r = results[0]
        assert r.file_path.endswith(".md")
        assert r.scope == "#"
        assert r.start_line >= 1
        assert r.end_line >= r.start_line

    def test_citation_metadata_for_code(self, indexed_engine):
        results = indexed_engine.search("verify_api_token", file_glob="*.py")
        assert results
        r = results[0]
        assert r.file_path.endswith(".py")
        assert r.start_line >= 1

    def test_empty_collection_returns_no_results(self, tmp_path, monkeypatch):
        empty = tmp_path / "empty"
        empty.mkdir()
        monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "data")
        engine = Holusight(empty, config=ServerConfig())
        engine.index()
        results = engine.search("anything")
        assert results == []

    def test_ask_without_llm_returns_graceful_or_mocked(self, indexed_engine):
        with patch("holusight.api.get_backend") as mock_get:
            backend = MagicMock()
            backend.model_id = "test:mock"
            backend.generate.return_value = "Net 30 applies. [Source 1]"
            mock_get.return_value = backend
            answer = indexed_engine.ask("What are the payment terms?")
        assert "Net 30" in answer.text or answer.sources
        assert answer.sources
        assert any("payment" in s.file_path for s in answer.sources)

    def test_index_persists_across_engine_restart(self, tmp_path, monkeypatch):
        data_dir = tmp_path / "data"
        monkeypatch.setattr(config_module, "DATA_DIR", data_dir)
        e1 = Holusight(FIXTURES, config=ServerConfig())
        e1.index(force_rebuild=True)
        count = e1.store.chunk_count
        e2 = Holusight(FIXTURES, config=ServerConfig())
        assert e2.store.is_indexed
        assert e2.store.chunk_count == count


@pytest.mark.parametrize("mode", ["ast-preamble", "ast-no-preamble", "regex-preamble"])
def test_index_search_emits_each_python_chunk_once(tmp_path, monkeypatch, mode):
    """Preamble chunks must not be embedded/stored twice or re-embedded unchanged."""
    if mode.startswith("ast"):
        pytest.importorskip("tree_sitter")
        pytest.importorskip("tree_sitter_python")
    else:
        def unavailable_parser(_language):
            raise ImportError("synthetic unavailable tree-sitter")

        monkeypatch.setattr("holusight.chunker._get_ts_parser", unavailable_parser)

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    preamble = "" if mode == "ast-no-preamble" else "import os\nimport sys\n\n"
    content = preamble + "def main():\n    return 'needle'\n"
    source = corpus / "main.py"
    source.write_text(content, encoding="utf-8")
    monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "index-data")

    class RecordingEmbedder:
        def __init__(self):
            self.texts = []

        def embed(self, texts):
            self.texts.extend(texts)
            return np.tile(np.array([1.0, 0.0], dtype=np.float32), (len(texts), 1))

        def embed_query(self, _query):
            return np.array([1.0, 0.0], dtype=np.float32)

    embedder = RecordingEmbedder()
    monkeypatch.setattr("holusight.indexer.get_embedder", lambda *a, **kw: embedder)
    monkeypatch.setattr("holusight.api.get_embedder", lambda *a, **kw: embedder)
    monkeypatch.setattr("holusight.indexer.VOYAGE_API_KEY", None)
    monkeypatch.setattr("holusight.search.VOYAGE_API_KEY", None)
    engine = Holusight(
        corpus,
        config=ServerConfig(
            embedding_model="synthetic",
            embedding_backend="local",
            embedding_dim=2,
            stale_threshold_seconds=3600,
            reranker=False,
            metadata_boost=False,
            query_enhancement=False,
            cnfb_alpha=0.0,
        ),
    )
    stores = []
    try:
        first = engine.index()
        results = engine.search("needle", top_k=8)
        stores.append(engine.store)
        vector_ids = engine.store.lance_table.to_arrow().column("chunk_id").to_pylist()
        expected_count = 1 if not preamble else 2

        assert results
        assert any("needle" in result.snippet for result in results)
        assert all(result.file_path == "main.py" for result in results)
        assert len(vector_ids) == len(set(vector_ids)) == expected_count
        assert first.chunks_created == first.total_chunks == expected_count
        assert len(embedder.texts) == len(set(embedder.texts)) == expected_count
        assert {result.chunk_id for result in results} == set(vector_ids)
        assert source.read_text(encoding="utf-8") == content
        assert list(corpus.iterdir()) == [source]

        original_embedding_texts = list(embedder.texts)
        second = engine.index()
        stores.append(engine.store)
        repeated = engine.search("needle", top_k=8)
        repeated_ids = engine.store.lance_table.to_arrow().column("chunk_id").to_pylist()
        assert second.chunks_created == 0
        assert second.chunks_skipped_unchanged == expected_count
        assert second.total_chunks == expected_count
        assert embedder.texts == original_embedding_texts
        assert repeated_ids == vector_ids
        assert {result.chunk_id for result in repeated} == set(vector_ids)
        assert source.read_text(encoding="utf-8") == content
    finally:
        for store in stores:
            store.close()
