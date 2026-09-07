"""Document parser regressions using real, generated temporary presentations."""

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from codesight.parsers import DocumentPage, extract_text


@pytest.mark.parametrize("placeholder", [True, False], ids=["placeholder", "textbox"])
def test_pptx_visible_text(tmp_path, caplog, placeholder):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1 if placeholder else 6])
    shape = (slide.placeholders[1] if placeholder else slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1),
    ))
    assert shape.is_placeholder is placeholder
    if not placeholder:
        # The real library property exists but is invalid for an ordinary textbox.
        with pytest.raises(ValueError, match="shape is not a placeholder"):
            _ = shape.placeholder_format
    shape.text = "  Textbox needle  \n\n  Second paragraph  "
    path = tmp_path / "visible.pptx"
    prs.save(path)
    original = path.read_bytes()

    assert extract_text(path) == [DocumentPage("Textbox needle\nSecond paragraph", 1)]
    assert "Failed to extract PPTX" not in caplog.text
    assert path.read_bytes() == original


@pytest.mark.parametrize("layout", [0, 1], ids=["center-title", "title"])
def test_pptx_title_metadata_without_duplicate_text(tmp_path, layout):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    slide.shapes.title.text = "  Slide title  "
    slide.placeholders[1].text = "Body text"
    assert all(shape.is_placeholder for shape in slide.shapes)
    path = tmp_path / "placeholder-only.pptx"
    prs.save(path)

    assert extract_text(path) == [DocumentPage("Slide title\nBody text", 1, "Slide title")]


@pytest.mark.parametrize("decoration", ["rectangle", "connector"])
def test_pptx_decoration_does_not_abort_remaining_slides(tmp_path, caplog, decoration):
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = "Opening title"
    prs.slides.add_slide(prs.slide_layouts[6])  # Empty slide still occupies slide number 2.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if decoration == "rectangle":
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(1),
        )
        assert shape.has_text_frame
    else:
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(0), Inches(0), Inches(1), Inches(1),
        )
        assert not shape.has_text_frame
    assert shape.is_placeholder is False
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = "Visible textbox"
    later = prs.slides.add_slide(prs.slide_layouts[1])
    later.shapes.title.text = "Later title"
    later.placeholders[1].text = "Later body"
    path = tmp_path / "mixed.pptx"
    prs.save(path)
    original = path.read_bytes()

    assert extract_text(path) == [
        DocumentPage("Opening title", 1, "Opening title"),
        DocumentPage("Visible textbox", 3),
        DocumentPage("Later title\nLater body", 4, "Later title"),
    ]
    assert "Failed to extract PPTX" not in caplog.text
    assert path.read_bytes() == original
