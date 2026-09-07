"""Real, generated DOCX regressions with local-only index/search controls."""

import numpy as np
import pytest
from docx import Document

from codesight import config as config_module
from codesight.api import CodeSight
from codesight.config import ServerConfig
from codesight.parsers import DocumentPage, extract_text
from codesight.store import ChunkStore


@pytest.fixture
def local_engine(tmp_path, monkeypatch):
    """Exercise real storage without models, providers, or shared data."""
    class FakeEmbedder:
        def embed(self, texts):
            return np.array([[1.0, 0.0] for _ in texts], dtype=np.float32)

        def embed_query(self, text):
            return np.array([1.0, 0.0], dtype=np.float32)

    stores = []

    def tracked_store(*args, **kwargs):
        store = ChunkStore(*args, **kwargs)
        stores.append(store)
        return store

    monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "data")
    for module in ("codesight.api", "codesight.indexer"):
        monkeypatch.setattr(f"{module}.get_embedder", lambda *a, **kw: FakeEmbedder())
        monkeypatch.setattr(f"{module}.ChunkStore", tracked_store)
    for module in ("codesight.indexer", "codesight.search"):
        monkeypatch.setattr(f"{module}.VOYAGE_API_KEY", None)
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    engine = CodeSight(corpus, config=ServerConfig(
        embedding_model="synthetic", embedding_backend="local", embedding_dim=2,
        reranker=False, metadata_boost=False,
    ))
    try:
        yield engine
    finally:
        for store in stores:
            store.close()


def test_table_only_index_search_with_paragraph_control(local_engine):
    engine = local_engine
    paragraph = Document()
    paragraph.add_paragraph("table needle")
    paragraph.save(engine.folder_path / "paragraph.docx")
    table = Document()
    table.add_table(rows=1, cols=1).cell(0, 0).text = "table needle"
    table.save(engine.folder_path / "table.docx")
    before = {p.name: p.read_bytes() for p in engine.folder_path.iterdir()}

    stats = engine.index()
    results = engine.search("table needle")
    hits = {r.file_path: r for r in results}
    assert "paragraph.docx" in hits, "Same-term paragraph is the causal control"
    assert "table.docx" in hits, "Body-table content must reach public search"
    assert stats.files_indexed == stats.total_chunks == 2
    assert set(engine.store.bm25_search("needle")) == {r.chunk_id for r in hits.values()}
    for result in hits.values():
        assert "table needle" in result.snippet
        assert result.start_line == result.end_line == 1
        assert result.scope == "page 1"
    assert {p.name: p.read_bytes() for p in engine.folder_path.iterdir()} == before


def test_body_order_and_heading_context_reach_search(local_engine):
    engine = local_engine
    doc = Document()
    doc.add_paragraph("Preface")
    doc.add_heading("Inventory", level=1)
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(
        [cell for row in table.rows for cell in row.cells],
        ["saffron needle", "quantity", "second row", "last cell"],
    ):
        cell.text = text
    doc.add_paragraph("After table")
    doc.add_heading("Appendix", level=2)
    doc.add_paragraph("Final paragraph")
    path = engine.folder_path / "ordered.docx"
    doc.save(path)

    expected = (
        "Inventory\nBefore table\nsaffron needle\nquantity\nsecond row\nlast cell\nAfter table"
    )
    assert extract_text(path) == [
        DocumentPage("Preface", 1),
        DocumentPage(expected, 2, "Inventory"),
        DocumentPage("Appendix\nFinal paragraph", 3, "Appendix"),
    ]
    engine.index()
    results = engine.search("saffron needle")
    assert results[0].snippet == expected
    assert results[0].scope == "Inventory"
    assert results[0].start_line == results[0].end_line == 2


@pytest.mark.parametrize("merge", ["horizontal", "vertical", "rectangle"])
def test_merged_cell_aliases_emit_once(tmp_path, merge):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    end_row, end_col = {"horizontal": (0, 1), "vertical": (1, 0), "rectangle": (1, 1)}[merge]
    table.cell(0, 0).merge(table.cell(end_row, end_col)).text = "merged needle"
    doc.add_paragraph("After table")
    path = tmp_path / "merged.docx"
    doc.save(path)
    assert extract_text(path) == [DocumentPage("merged needle\nAfter table", 1)]


def test_distinct_cells_with_equal_text_are_not_deduplicated(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    for cell in table.rows[0].cells:
        cell.text = "same text"
    path = tmp_path / "equal.docx"
    doc.save(path)
    assert extract_text(path) == [DocumentPage("same text\nsame text", 1)]


def test_nested_table_and_cell_paragraph_order(tmp_path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Cell before"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested needle"
    cell.add_paragraph("Cell after")
    doc.add_paragraph("Body after")
    path = tmp_path / "nested.docx"
    doc.save(path)
    assert extract_text(path) == [
        DocumentPage("Cell before\nNested needle\nCell after\nBody after", 1),
    ]


def test_paragraph_only_sections_and_whitespace_unchanged(tmp_path):
    doc = Document()
    doc.add_paragraph("  Preface  ")
    doc.add_paragraph("   ")
    doc.add_heading(" First ", level=1)
    doc.add_paragraph("  Body  ")
    doc.add_heading("Second", level=2)
    doc.add_heading("Third", level=1)
    doc.add_paragraph("End")
    path = tmp_path / "paragraphs.docx"
    doc.save(path)
    assert extract_text(path) == [
        DocumentPage("Preface", 1),
        DocumentPage("First \n  Body", 2, "First"),
        DocumentPage("Second", 3, "Second"),
        DocumentPage("Third\nEnd", 4, "Third"),
    ]


def test_empty_body_ignores_headers_and_footers(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Header excluded"
    doc.sections[0].footer.paragraphs[0].text = "Footer excluded"
    doc.add_table(rows=1, cols=1)
    path = tmp_path / "empty.docx"
    doc.save(path)
    assert extract_text(path) == []
