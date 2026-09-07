"""Real synthetic DOCX indexing and retrieval, with no model or provider calls."""

import numpy as np
import pytest
from docx import Document

from codesight import config as config_module
from codesight.api import CodeSight
from codesight.chunker import chunk_document
from codesight.config import ServerConfig
from codesight.parsers import extract_text
from codesight.store import ChunkStore


@pytest.mark.parametrize("layout", ["oversized", "short-paragraphs", "near-limit-overlap"])
def test_document_index_bounds_and_tail_citation(tmp_path, monkeypatch, layout):
    monkeypatch.setenv("HF_HUB_OFFLINE", "1")
    monkeypatch.setenv("TRANSFORMERS_OFFLINE", "1")
    monkeypatch.setattr(config_module, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr("codesight.indexer.VOYAGE_API_KEY", None)
    monkeypatch.setattr("codesight.search.VOYAGE_API_KEY", None)
    embedded = []

    class FakeEmbedder:
        def embed(self, texts):
            embedded.extend(texts)
            return np.array([[1.0, 0.0]] * len(texts), dtype=np.float32)

        def embed_query(self, query):
            return np.array([1.0, 0.0], dtype=np.float32)

    fake = FakeEmbedder()
    monkeypatch.setattr("codesight.api.get_embedder", lambda *a, **kw: fake)
    monkeypatch.setattr("codesight.indexer.get_embedder", lambda *a, **kw: fake)
    # Retain and close all test-owned stores, including the indexer's local store.
    stores = []

    def make_store(*args, **kwargs):
        store = ChunkStore(*args, **kwargs)
        stores.append(store)
        return store

    monkeypatch.setattr("codesight.api.ChunkStore", make_store)
    monkeypatch.setattr("codesight.indexer.ChunkStore", make_store)
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    path = corpus / "sample.docx"
    tail = "quartzfinchuniquetail"
    parts = [f"Synthetic item {i:03d} contains public demonstration text." for i in range(40)]
    parts.append(tail)
    doc = Document()
    if layout == "oversized":
        doc.add_paragraph(" ".join(parts))
    elif layout == "short-paragraphs":
        # Blank lines preserve paragraph boundaries in the extracted DOCX section.
        for part in parts:
            doc.add_paragraph(part + "\n")
    else:
        doc.add_paragraph("a" * 180 + "\n")
        doc.add_paragraph("b" * 179 + " " + tail)
    doc.save(path)
    source_bytes = path.read_bytes()
    config = ServerConfig(
        embedding_model="synthetic", embedding_backend="local", embedding_dim=2,
        doc_chunk_max_chars=200, doc_chunk_overlap_chars=20,
        reranker=False, metadata_boost=False,
    )
    engine = CodeSight(corpus, config=config)
    try:
        stats = engine.index()
        assert stats.files_indexed == 1
        ids = list(engine.store.fts.get_chunk_hashes("sample.docx"))
        rows = engine.store.get_chunk_metadata(ids)
        assert rows
        assert all(0 < len(row["content"]) <= 200 for row in rows.values())
        expected = chunk_document(extract_text(path), "sample.docx", 200, 20)
        assert set(ids) == {c.chunk_id for c in expected}
        assert stats.chunks_created == stats.total_chunks == len(expected)
        assert embedded == [c.embedding_text for c in expected]
        for chunk in expected:
            row = rows[chunk.chunk_id]
            assert row["content"] == chunk.content
            assert row["content_hash"] == chunk.content_hash
            assert (row["start_line"], row["end_line"], row["scope"]) == (1, 1, "page 1")
            assert row["language"] == "docx"
        assert engine.store.bm25_search(tail)
        hits = engine.search(tail, file_glob="*.docx")
        tail_hits = [hit for hit in hits if tail in hit.snippet]
        assert tail_hits
        assert all((h.file_path, h.start_line, h.end_line) == ("sample.docx", 1, 1)
                   for h in tail_hits)
        embedded.clear()
        second = engine.index()
        assert second.chunks_created == 0
        assert second.chunks_skipped_unchanged == len(expected)
        assert not embedded
        assert engine.store.get_chunk_metadata(ids) == rows
        assert path.read_bytes() == source_bytes
        assert list(corpus.iterdir()) == [path]
    finally:
        for store in stores:
            store.close()
